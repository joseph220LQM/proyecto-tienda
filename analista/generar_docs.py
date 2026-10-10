from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

AZUL = RGBColor(0x1F, 0x49, 0x7D)
GRIS = RGBColor(0x44, 0x47, 0x4A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
AZUL_CLARO = RGBColor(0xBD, 0xD7, 0xEE)


def set_cell_bg(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = AZUL
    return p


def add_table(doc, headers, rows, header_color="1F497D"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_bg(hdr[i], header_color)
        run = hdr[i].paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = BLANCO
        run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            cells[ci].paragraphs[0].runs[0].font.size = Pt(9)
            if ri % 2 == 1:
                set_cell_bg(cells[ci], "EBF3FB")
    return table


# ── DOCUMENTO 1: REQUERIMIENTOS ──────────────────────────────────────────────
doc1 = Document()

# Márgenes
for section in doc1.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

# Portada
p = doc1.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("DOCUMENTO DE REQUERIMIENTOS")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = AZUL

p2 = doc1.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Sistema de Gestión de Ventas\nArtículos Masculinos y Femeninos")
r2.font.size = Pt(13)
r2.font.color.rgb = GRIS

doc1.add_paragraph()
meta = doc1.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Proyecto: Tienda de Artículos  |  Área: Creatividad y Desarrollo  |  Octubre 2026").font.size = Pt(10)

doc1.add_page_break()

# 1. Introducción
add_heading(doc1, "1. Introducción")
doc1.add_paragraph(
    "Este documento describe los requerimientos del sistema de escritorio para la gestión de clientes, "
    "productos y ventas de una tienda de artículos masculinos y femeninos. El sistema busca reemplazar "
    "los procesos manuales actuales con una solución integrada que facilite el control del inventario "
    "y el seguimiento de las operaciones comerciales."
)

# 2. Actores
add_heading(doc1, "2. Actores del sistema")
add_table(doc1,
    ["Actor", "Descripción"],
    [
        ["Vendedor", "Actor principal. Gestiona clientes, registra ventas y consulta productos."],
        ["Administrador", "Actor secundario. Consulta reportes, revisa ventas y verifica información registrada."],
    ]
)

# 3. RF
doc1.add_paragraph()
add_heading(doc1, "3. Requerimientos Funcionales")

add_heading(doc1, "Módulo de Clientes", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RF01", "Registrar cliente", "Permite registrar un nuevo cliente con: identificación, nombre, teléfono, correo y dirección. No se permiten identificaciones duplicadas."],
        ["RF02", "Consultar clientes", "Muestra la lista de clientes en una tabla y permite buscar por nombre o identificación."],
        ["RF03", "Modificar cliente", "Permite seleccionar un cliente y actualizar sus datos registrados."],
        ["RF04", "Eliminar cliente", "Permite eliminar un cliente aplicando validaciones para evitar eliminar clientes con ventas asociadas."],
    ]
)

doc1.add_paragraph()
add_heading(doc1, "Módulo de Productos", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RF05", "Registrar producto", "Permite registrar: código único, nombre, descripción, precio, categoría, género (masculino/femenino), talla y cantidad disponible."],
        ["RF06", "Consultar productos", "Muestra todos los productos en una tabla y permite filtrar por nombre, categoría o género."],
        ["RF07", "Modificar producto", "Permite seleccionar un producto y actualizar sus datos, validando los campos antes de guardar."],
        ["RF08", "Eliminar producto", "Permite eliminar un producto validando que no esté asociado a ventas registradas."],
    ]
)

doc1.add_paragraph()
add_heading(doc1, "Módulo de Ventas", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RF09", "Registrar venta", "Permite crear una venta seleccionando un cliente y uno o varios productos con su cantidad. El total se calcula automáticamente."],
        ["RF10", "Consultar ventas", "Muestra el historial de ventas con fecha, cliente y total."],
        ["RF11", "Actualizar inventario", "Al registrar una venta exitosa, descuenta automáticamente la cantidad vendida del inventario de cada producto."],
        ["RF12", "Consultar detalle de venta", "Muestra el detalle completo de una venta: cliente, productos, cantidades, precios unitarios y total."],
    ]
)

# 4. RNF
doc1.add_paragraph()
add_heading(doc1, "4. Requerimientos No Funcionales")

add_heading(doc1, "Plataforma y tecnología", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RNF01", "Tipo de aplicación", "La aplicación será de escritorio, diseñada para ejecutarse en Windows."],
        ["RNF02", "Interfaz gráfica", "La interfaz de usuario se desarrollará con Java Swing."],
        ["RNF03", "Base de datos", "El sistema utilizará MySQL como motor de base de datos."],
        ["RNF04", "Conexión BD", "La comunicación se realizará mediante JDBC, usando el conector MySQL versión 9.6 o 9.7."],
    ]
)

doc1.add_paragraph()
add_heading(doc1, "Validaciones y comportamiento", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RNF05", "Campos obligatorios", "Todos los formularios deben validar que los campos requeridos estén diligenciados antes de procesar cualquier operación."],
        ["RNF06", "Valores numéricos", "Los campos de precio y cantidad deben aceptar únicamente valores mayores que cero."],
        ["RNF07", "Control de inventario", "No se debe permitir registrar una venta si la cantidad solicitada supera las unidades disponibles."],
        ["RNF08", "Mensajes al usuario", "El sistema debe mostrar mensajes claros ante operaciones exitosas, errores de validación o fallos de conexión."],
    ]
)

doc1.add_paragraph()
add_heading(doc1, "Mantenibilidad", level=2)
add_table(doc1,
    ["Código", "Requerimiento", "Descripción"],
    [
        ["RNF09", "Organización del código", "El código debe estructurarse por responsabilidades (conexión, lógica de negocio, interfaz) para facilitar su mantenimiento."],
    ]
)

# 5. RN
doc1.add_paragraph()
add_heading(doc1, "5. Reglas del negocio")
add_table(doc1,
    ["Código", "Regla"],
    [
        ["RN01", "Un cliente no puede tener dos registros con la misma identificación."],
        ["RN02", "Un producto debe tener un código único en el sistema."],
        ["RN03", "Todo producto debe estar clasificado como masculino o femenino."],
        ["RN04", "El precio de cualquier producto debe ser mayor que cero."],
        ["RN05", "La cantidad disponible de un producto no puede ser negativa."],
        ["RN06", "Toda venta debe estar asociada obligatoriamente a un cliente registrado."],
        ["RN07", "Una venta debe contener como mínimo un producto."],
        ["RN08", "La cantidad de cada producto vendido debe ser mayor que cero."],
        ["RN09", "No se puede registrar una venta si la cantidad solicitada supera el inventario disponible."],
        ["RN10", "El total de la venta se calcula automáticamente: suma de (precio × cantidad) por cada producto."],
        ["RN11", "Una venta válida queda registrada de forma definitiva en el sistema."],
        ["RN12", "El inventario se actualiza automáticamente al completarse una venta."],
    ]
)

doc1.save("C:/Users/user/Documents/MozartBPO/Proyecto1_Tienda/analista/Requerimientos.docx")
print("Requerimientos.docx generado")


# ── DOCUMENTO 2: HISTORIAS DE USUARIO ────────────────────────────────────────
doc2 = Document()

for section in doc2.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

p = doc2.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("HISTORIAS DE USUARIO — BACKLOG")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = AZUL

p2 = doc2.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Sistema de Gestión de Ventas\nArtículos Masculinos y Femeninos")
r2.font.size = Pt(13)
r2.font.color.rgb = GRIS

doc2.add_paragraph()
meta = doc2.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Proyecto: Tienda de Artículos  |  Área: Creatividad y Desarrollo  |  Octubre 2026").font.size = Pt(10)

doc2.add_paragraph()
nota = doc2.add_paragraph(
    "Cada historia de usuario corresponde a un Issue en GitHub, organizado en el tablero "
    "GitHub Projects con columnas: Por hacer → En progreso → Hecho."
)
nota.runs[0].font.size = Pt(10)
nota.runs[0].font.italic = True

doc2.add_page_break()

historias = [
    {
        "id": "HU01", "nombre": "Registrar cliente",
        "etiqueta": "feature  |  módulo-clientes",
        "historia": "Como vendedor, quiero registrar un nuevo cliente en el sistema para poder asociarlo a sus compras futuras.",
        "descripcion": "El formulario de registro debe incluir: número de identificación, nombre completo, teléfono, correo electrónico y dirección.",
        "restricciones": [
            "El número de identificación es único; no se admiten duplicados.",
            "El correo debe tener formato válido.",
        ],
        "criterios": [
            "Muestra formulario con todos los campos requeridos.",
            "No permite guardar si algún campo obligatorio está vacío.",
            "Muestra error si la identificación ya existe.",
            "Confirma con mensaje cuando el cliente se registra exitosamente.",
        ],
    },
    {
        "id": "HU02", "nombre": "Consultar cliente",
        "etiqueta": "feature  |  módulo-clientes",
        "historia": "Como vendedor, quiero consultar los clientes registrados para encontrar rápidamente su información.",
        "descripcion": "El sistema muestra todos los clientes en una tabla con campo de búsqueda por nombre o identificación.",
        "restricciones": ["La tabla debe actualizarse al realizar una búsqueda."],
        "criterios": [
            "Tabla con columnas: ID, nombre, teléfono, correo y dirección.",
            "Campo de búsqueda filtra por nombre o identificación.",
            "Si no hay resultados, muestra un mensaje indicándolo.",
        ],
    },
    {
        "id": "HU03", "nombre": "Registrar producto",
        "etiqueta": "feature  |  módulo-productos",
        "historia": "Como encargado, quiero registrar nuevos productos para mantener actualizado el catálogo de artículos disponibles.",
        "descripcion": "Formulario con campos: código, nombre, descripción, precio, categoría, género (masculino/femenino), talla y cantidad disponible.",
        "restricciones": [
            "El código del producto debe ser único.",
            "El precio debe ser mayor que cero.",
            "El género solo puede ser 'Masculino' o 'Femenino'.",
        ],
        "criterios": [
            "Formulario con todos los campos requeridos.",
            "No permite guardar con precio o cantidad inválidos.",
            "Muestra error si el código ya existe.",
            "Confirma el registro exitoso con un mensaje.",
        ],
    },
    {
        "id": "HU04", "nombre": "Consultar producto",
        "etiqueta": "feature  |  módulo-productos",
        "historia": "Como vendedor, quiero consultar los productos disponibles para conocer su precio y cantidad en inventario.",
        "descripcion": "Tabla con todos los productos, filtrable por nombre, categoría o género.",
        "restricciones": ["Debe mostrar claramente la cantidad disponible de cada producto."],
        "criterios": [
            "Tabla con: código, nombre, categoría, género, talla, precio y cantidad disponible.",
            "Permite filtrar por género o categoría.",
            "Productos con cantidad cero se diferencian visualmente.",
        ],
    },
    {
        "id": "HU05", "nombre": "Registrar venta",
        "etiqueta": "feature  |  módulo-ventas",
        "historia": "Como vendedor, quiero registrar una nueva venta para almacenar los productos adquiridos por un cliente.",
        "descripcion": "Selección de cliente, agregado de productos con cantidad, cálculo automático del total y confirmación.",
        "restricciones": [
            "No se puede registrar sin cliente seleccionado.",
            "Debe haber al menos un producto.",
            "No se puede vender más unidades de las disponibles.",
            "El total es de solo lectura; se calcula automáticamente.",
        ],
        "criterios": [
            "Permite seleccionar cliente desde la lista.",
            "Permite agregar múltiples productos con cantidad editable.",
            "El total se calcula y actualiza automáticamente.",
            "Muestra error si se supera el stock disponible.",
            "Al confirmar, la venta queda registrada y el inventario se descuenta.",
        ],
    },
    {
        "id": "HU06", "nombre": "Actualizar inventario tras venta",
        "etiqueta": "feature  |  módulo-ventas  |  módulo-productos",
        "historia": "Como encargado, quiero que el inventario se actualice automáticamente después de cada venta.",
        "descripcion": "Al confirmar la venta, el sistema resta la cantidad vendida al stock de cada producto involucrado.",
        "restricciones": [
            "El descuento ocurre solo si la venta se registra correctamente.",
            "Si la venta falla, el inventario no debe modificarse.",
        ],
        "criterios": [
            "La cantidad disponible disminuye según lo vendido.",
            "El cambio se refleja inmediatamente al consultar productos.",
            "Si la venta falla, el inventario permanece sin cambios.",
        ],
    },
    {
        "id": "HU07", "nombre": "Consultar ventas",
        "etiqueta": "feature  |  módulo-ventas",
        "historia": "Como vendedor, quiero consultar el historial de ventas para conocer las operaciones realizadas.",
        "descripcion": "Tabla con todas las ventas, ordenadas de más reciente a más antigua.",
        "restricciones": ["Ordenadas por fecha descendente."],
        "criterios": [
            "Tabla con: número de venta, fecha, nombre del cliente y total.",
            "Al seleccionar una venta se puede acceder a su detalle.",
        ],
    },
    {
        "id": "HU08", "nombre": "Modificar producto",
        "etiqueta": "feature  |  módulo-productos",
        "historia": "Como encargado, quiero modificar un producto para mantener actualizados sus datos.",
        "descripcion": "Al seleccionar un producto se cargan sus datos en un formulario editable para guardar cambios.",
        "restricciones": ["Se deben validar los datos antes de guardar."],
        "criterios": [
            "Los datos del producto se cargan en el formulario al seleccionarlo.",
            "Permite editar y guardar con validación de campos.",
            "Muestra confirmación al guardar exitosamente.",
            "Muestra error si los datos ingresados no son válidos.",
        ],
    },
    {
        "id": "HU09", "nombre": "Consultar detalle de venta",
        "etiqueta": "feature  |  módulo-ventas",
        "historia": "Como administrador, quiero consultar el detalle de una venta para verificar los productos vendidos y el total.",
        "descripcion": "Vista de solo lectura con información completa de la venta seleccionada.",
        "restricciones": ["El detalle es de solo lectura; una venta registrada no puede modificarse."],
        "criterios": [
            "Muestra datos del cliente asociado.",
            "Lista productos con cantidad y precio unitario.",
            "Muestra el total de la venta.",
            "Incluye fecha y número de la venta.",
        ],
    },
]

for hu in historias:
    h = add_heading(doc2, f"{hu['id']} — {hu['nombre']}")

    etiq = doc2.add_paragraph()
    r = etiq.add_run(f"Etiqueta: {hu['etiqueta']}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    r.bold = True

    doc2.add_paragraph().add_run("Historia:").bold = True
    doc2.add_paragraph(hu["historia"]).runs[0].font.size = Pt(10)

    doc2.add_paragraph().add_run("Descripción:").bold = True
    doc2.add_paragraph(hu["descripcion"]).runs[0].font.size = Pt(10)

    doc2.add_paragraph().add_run("Restricciones:").bold = True
    for r in hu["restricciones"]:
        p = doc2.add_paragraph(r, style="List Bullet")
        p.runs[0].font.size = Pt(10)

    doc2.add_paragraph().add_run("Criterios de aceptación:").bold = True
    for c in hu["criterios"]:
        p = doc2.add_paragraph(f"☐  {c}", style="List Bullet")
        p.runs[0].font.size = Pt(10)

    doc2.add_paragraph()

# Tabla resumen backlog
add_heading(doc2, "Organización sugerida en GitHub Projects")
add_table(doc2,
    ["Columna", "Issues"],
    [
        ["Por hacer", "HU01, HU02, HU03, HU04, HU05, HU06, HU07, HU08, HU09"],
        ["En progreso", "(se mueven según avance del equipo)"],
        ["Hecho", "(se mueven al completar y validar)"],
    ]
)

doc2.save("C:/Users/user/Documents/MozartBPO/Proyecto1_Tienda/analista/HistoriasDeUsuario.docx")
print("HistoriasDeUsuario.docx generado")
